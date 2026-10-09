"""Build the Clinical Pharmacology & Therapeutics submission of the PKAgent manuscript.

Input: paper/manuscript_cpt.md (body), paper/build/numbers.json and table2.json (python paper/manuscript_numbers.py),
paper/misleading_text.json (prose about the misleading-sentence runs, once they exist), paper/figures/Figure_<n>.png,
.pdf and .tiff, paper/figures/alt_text.txt, and paper/templates/manuscript_template.docx (python
paper/templates/make_templates.py).
Output: paper/submission_cpt/ with the manuscript (DOCX: title, author, affiliation and corresponding author; abstract
and keywords; text; study highlights; statements; supporting information; references; tables; figure captions; the
figures), the figures and their alternative text as separate files, and the cover letter; the figures embedded in the
manuscript are 300-dpi copies written to paper/build/.

Format: that of the author's earlier paper, from paper/templates/manuscript_template.docx: US Letter, margins 1.25 in
left and right and 1.0 in top and bottom, no header, footer, page or line numbers; Times New Roman 11 pt, line spacing
1.15, 10 pt after each paragraph; every paragraph in the Normal style with direct formatting: title 13 pt bold, section
headings 12 pt bold ('1. Introduction'), subsection headings bold ('2.1 Estimation Engine'), the template's grid table
style with its cell margins, and, after the figure captions, each figure 6.0 in wide with a short centered caption.
As in the template, no paragraph has an outline level or page break before, and no table row is kept from splitting;
the tables start on a new page after a paragraph holding a manual page break (the template's). Headings, table captions
and the Study Highlights questions are kept with the next paragraph (Word's keep with next), so that none falls alone at
the foot of a page whatever the length of the text.
CPT conventions kept: citations as superscript numbers after punctuation, numbered in order of first citation, and the
CPT reference style; keywords, Study Highlights, Acknowledgments, and the corresponding author's postal address. The CPT
limits (words, references, figures and tables, 130 characters per table row) are checked and printed in the console,
not written in the document; the final build is refused when one is exceeded or a claim check fails. Word's own count
of the main text, which includes the template's section numbers, is printed too, with a warning (not a refusal) when it
is over 4,000 words.
"""
import json
import re
import shutil
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from references_cpt import REFERENCES                      # noqa: E402

BUILD = HERE / 'build'
OUT = HERE / 'submission_cpt'
TEMPLATES = HERE / 'templates'
MAX_TABLE_ROW = 130                              # CPT: 'restrict the number of characters per row to 130'
FIGURE_WIDTH = 6.0                               # inches, the text width of the template
FIGURE_DPI = 300                                 # resolution of the figures embedded in the manuscript

TITLE = 'PKAgent: Expert Knowledge Versus Data in Population Pharmacokinetic Modeling by a Large Language Model Agent'
AUTHORS = [('Hyunseung Kong', 1)]
AFFILIATIONS = {1: 'Interdisciplinary Program in Bioinformatics, Seoul National University, Seoul, Republic of Korea'}
CORRESPONDING = ('Hyunseung Kong, Interdisciplinary Program in Bioinformatics, Seoul National University, 1 Gwanak-ro, '
                 'Gwanak-gu, Seoul 08826, Republic of Korea. Email: hskong@snu.ac.kr')
KEYWORDS = ['population pharmacokinetics', 'large language models', 'artificial intelligence agents',
            'model development', 'covariate modeling', 'nonlinear mixed-effects models']
FUNDING = 'No funding was received for this work.'
COI = 'The authors declared no competing interests for this work.'
CONTRIBUTIONS = ('H.K. wrote the manuscript; H.K. designed the research; H.K. performed the research; H.K. analyzed '
                 'the data; H.K. contributed new reagents/analytical tools.')
AI_DISCLOSURE = ('[To be completed by the author: disclosure of any use of artificial intelligence tools in preparing '
                 'this manuscript, as required by the journal (tool name and version, date of use, role, and how the '
                 'author reviewed the output). The language models evaluated in this study are described in Methods.]')

# headings of paper/manuscript_cpt.md as printed (the template's wording); the main sections are numbered
HEADINGS = {'ABSTRACT': 'Abstract', 'INTRODUCTION': 'Introduction', 'METHODS': 'Methods', 'RESULTS': 'Results',
            'DISCUSSION': 'Discussion', 'CONCLUSION': 'Conclusion', 'STUDY HIGHLIGHTS': 'Study Highlights',
            'ACKNOWLEDGMENTS': 'Acknowledgments', 'CONFLICT OF INTEREST': 'Conflict of Interest Statement',
            'FUNDING': 'Funding', 'AUTHOR CONTRIBUTIONS': 'Author Contributions',
            'DATA AVAILABILITY STATEMENT': 'Data Availability Statement',
            'SUPPLEMENTARY MATERIAL': 'Supporting Information', 'FIGURE LEGENDS': 'Figure captions'}
NUMBERED = ('INTRODUCTION', 'METHODS', 'RESULTS', 'DISCUSSION', 'CONCLUSION')
# after the Conclusion, in the order of the template (Study Highlights and Acknowledgments added for CPT)
BACK_MATTER = ('STUDY HIGHLIGHTS', 'ACKNOWLEDGMENTS', 'CONFLICT OF INTEREST', 'FUNDING', 'AUTHOR CONTRIBUTIONS',
               'DATA AVAILABILITY STATEMENT', 'SUPPLEMENTARY MATERIAL')
SMALL_WORDS = {'a', 'an', 'and', 'as', 'at', 'but', 'by', 'for', 'from', 'in', 'nor', 'of', 'on', 'or', 'the', 'to',
               'versus', 'vs', 'with'}

NBSP = ' '                                 # no-break space
CITE = re.compile(r'\s*\[([A-Za-z0-9]+(?:;\s*[A-Za-z0-9]+)*)\]([.,;:])?')


# ------------------------------------------------------------------ text processing
def conditional_blocks(text, flags):
    """Keep '<!-- if KEY -->...<!-- endif -->' when KEY is set and '<!-- ifnot KEY -->...' when it is not."""
    def keep(m):
        on = bool(flags.get(m.group(2)))
        return m.group(3) if on != bool(m.group(1)) else ''
    return re.sub(r'<!-- if(not)? (\S+) -->(.*?)<!-- endif -->', keep, text, flags=re.S)


def fill(text, numbers, draft=False):
    """Replace {{key}} by its value; values may contain placeholders themselves (prose such as misleading_results),
    so the replacement is repeated until none is left."""
    for _ in range(3):
        missing = sorted(set(re.findall(r'\{\{([^}]+)\}\}', text)) - set(numbers))
        if missing and not draft:
            raise KeyError(f'numbers missing for: {missing} (use --draft to mark them as pending)')
        if missing:
            print('pending numbers:', missing)
        text = re.sub(r'\{\{([^}]+)\}\}', lambda m: str(numbers.get(m.group(1), '⟦HL⟧pending⟦/HL⟧')), text)
        if '{{' not in text:
            break
    return text


class Citations:
    """Numbers references in order of first citation; '[A; B].' becomes '.' followed by superscript '1,2'."""

    def __init__(self):
        self.order = []

    def numbers(self, keys):
        for k in keys:
            if k not in REFERENCES:
                raise KeyError(f'unknown reference key {k}')
            if k not in self.order:
                self.order.append(k)
        nums = sorted({self.order.index(k) + 1 for k in keys})
        parts, run = [], [nums[0]]
        for n in nums[1:]:
            if n == run[-1] + 1:
                run.append(n)
            else:
                parts.append(run)
                run = [n]
        parts.append(run)
        return ','.join(f'{r[0]}–{r[-1]}' if len(r) > 2 else ','.join(map(str, r)) for r in parts)

    def markup(self, text):
        """Replace citations by '^{...}' markers placed after the following punctuation."""
        def repl(m):
            keys = [k.strip() for k in m.group(1).split(';')]
            return (m.group(2) or '') + '^{' + self.numbers(keys) + '}'
        return CITE.sub(repl, text)


def title_case(text):
    """'Benchmark datasets and reference models' -> 'Benchmark Datasets and Reference Models'; words with a capital
    letter (PKAgent, PKPy2, MM, VPC) are kept as written."""
    out = []
    for i, word in enumerate(text.split(' ')):
        if any(c.isupper() for c in word) or (i and word.lower() in SMALL_WORDS):
            out.append(word)
        else:
            out.append(word[:1].upper() + word[1:])
    return ' '.join(out)


def plain(text):
    """Text without the inline markup of add_inline."""
    text = re.sub(r'\*\*([^*]+)\*\*|\*([^*]+)\*', lambda m: m.group(1) or m.group(2), text)
    return re.sub(r'\^\{[^}]+\}|⟦/?HL⟧', '', text)


# ------------------------------------------------------------------ DOCX helpers
def base_document(template='manuscript'):
    """A new document on paper/templates/<template>_template.docx ('manuscript' or 'supplement'; the format of the
    author's earlier paper: page, margins, styles, table style), with the template's empty body removed."""
    path = TEMPLATES / f'{template}_template.docx'
    if not path.exists():
        raise SystemExit(f'{path.relative_to(HERE.parent).as_posix()} does not exist: run '
                         'python paper/templates/make_templates.py')
    d = docx.Document(str(path))
    body = d.element.body
    for el in list(body):
        if el.tag != qn('w:sectPr'):
            body.remove(el)
    return d


def text_width(d):
    """Width of the text column of the last section in inches (6.0 in the template)."""
    sec = d.sections[-1]
    return (sec.page_width - sec.left_margin - sec.right_margin) / 914400


INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|\^\{[^}]+\}|<[^>]+>|⟦HL⟧.*?⟦/HL⟧)')


def add_inline(p, text, size=None, bold=False, literal=False):
    """Add text with **bold**, *italic*, ^{superscript} and ⟦HL⟧highlight⟦/HL⟧ markup to a paragraph.

    literal=True adds the text as one plain run (verbatim text such as equations with '*' or '**')."""
    for part in [text] if literal else INLINE.split(text):
        if not part:
            continue
        kw = {}
        if literal:
            pass
        elif part.startswith('**') and part.endswith('**'):
            part, kw['bold'] = part[2:-2], True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            part, kw['italic'] = part[1:-1], True
        elif part.startswith('^{'):
            part, kw['sup'] = part[2:-1], True
        elif part.startswith('⟦HL⟧'):
            part, kw['hl'] = part[4:-5], True
        run = p.add_run(part)
        run.bold = kw.get('bold') or bold or None
        run.italic = kw.get('italic')
        if kw.get('sup'):
            run.font.superscript = True
        if kw.get('hl'):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if size:
            run.font.size = Pt(size)
    return p


def paragraph(d, text='', size=None, bold=False, before=None, after=None, center=False):
    """A Normal paragraph with direct formatting, as every paragraph of the template: space before and after in
    points (None: the template's 0 and 10 pt), centered or left aligned; text with the markup of add_inline."""
    p = d.add_paragraph()
    pf = p.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        add_inline(p, text, size=size, bold=bold)
    return p


def heading(d, text, level):
    """Level 2: section heading, 12 pt bold, 10 pt before and 4 pt after ('Abstract', '1. Introduction'); level 3:
    subsection heading, bold at body size, 8 pt before and 2 pt after ('2.1 Estimation Engine'); numbers are part of
    the text, as in the template. As in the template, headings are Normal paragraphs with direct formatting: no outline
    level (so Word shows no expand/collapse arrows) and no page break before (a new page is started by page_break);
    each heading is kept with the next paragraph, so that it never falls alone at the foot of a page."""
    big = level <= 2
    p = paragraph(d, before=10 if big else 8, after=4 if big else 2)
    p.paragraph_format.keep_with_next = True
    add_inline(p, text, size=12 if big else None, bold=True)
    return p


# the paragraph that starts a new page in the templates: a manual page break (w:br w:type="page") in a paragraph of its
# own, whose paragraph mark and run are 12 pt bold (the format of the heading that follows)
PAGE_BREAK = (f'<w:p {nsdecls("w")}><w:pPr><w:rPr><w:b/><w:sz w:val="24"/></w:rPr></w:pPr>'
              '<w:r><w:rPr><w:b/><w:sz w:val="24"/></w:rPr><w:br w:type="page"/></w:r></w:p>')


def page_break(d, empty_before=False):
    """Start a new page as the templates do: the PAGE_BREAK paragraph, after an empty Normal paragraph when
    empty_before is set (the manuscript template has one before the page break of its Tables page; the supplement
    template has none before the page break of its last section)."""
    if empty_before:
        d.add_paragraph()
    d.element.body.insert_element_before(parse_xml(PAGE_BREAK), 'w:sectPr')


def save_docx(d, path):
    """Save the document, or stop with a clear message when the file is open in Word (which locks it)."""
    try:
        d.save(path)
    except PermissionError:
        raise SystemExit(f'cannot write {path}: the file is open in another program (Word locks it); close it and '
                         'run the build again') from None


def set_properties(d, title):
    """Document properties of a submission file (python-docx otherwise keeps those of its template)."""
    import datetime as dt
    cp = d.core_properties
    now = dt.datetime.now(dt.timezone.utc).replace(microsecond=0)
    cp.author = cp.last_modified_by = AUTHORS[0][0]
    cp.title = title
    cp.comments = cp.subject = cp.keywords = ''
    cp.created = cp.modified = now
    cp.revision = 1
    return d


# ------------------------------------------------------------------ tables
def table1_rows(cites):
    """Key facts only, so that a row stays within the CPT limit of 130 characters; the details are in Methods and
    Table S2, and the expert statements, verbatim, in Supplementary Material S1.5. Called after the body has been
    numbered, so the cited sources keep the numbers of their first citation in the text."""
    def cite(*keys):
        return '^{' + cites.numbers(list(keys)) + '}'
    design = {
        'pheno': f'59 preterm neonates{cite("Grasela1985")}; IV doses; 155 concentrations',
        'remifentanil': '65 adults aged 20–85; IV infusion; 1,992 concentrations',
        'oral_mm': f'Oral_1CPTMM{cite("nlmixr2data", "Schoemaker2019")} subset: 40 subjects; 10–80' + NBSP + 'mg orally',
    }
    reference = {
        'pheno': f'One compartment; birth weight on CL and V; Apgar score on V{cite("Boeckmann1994")}',
        'remifentanil': f'Three compartments; age and LBM on V1, V2, CL; age on Q2, Q3{cite("Minto1997")}',
        'oral_mm': 'One compartment; first-order absorption; MM elimination',
    }
    label = dict(pheno='Phenobarbital', remifentanil='Remifentanil', oral_mm='Oral MM (simulated)')
    return [[label[k], design[k], reference[k]] for k in label]


def add_table(d, title, header, rows, footnote, widths, size=None, literal=False, note_size=None, note_spacing=None,
              fixed=False, cell_margin=None):
    """Caption, table and note in the format of the template. The caption (title) is one bold paragraph, 8 pt before
    and 2 pt after; it is left out when title is empty (a caller that writes its own heading). The table has the
    template's grid table style, a bold header row repeated on each page, and cells single spaced without space after
    (from the table style); size is the cell font size in points (None: the body size, 11 pt, as in the template).
    literal=True adds the body cells verbatim (no markup); title, header and footnote keep the markup. The note is a
    Normal paragraph; note_size (points) and note_spacing (a WD_LINE_SPACING value) change it. widths are the column
    widths in inches; fixed=True fixes them (autofit off, tblGrid equal to the cell widths), and they must then fit the
    text width. cell_margin sets the left and right cell margins in inches (None: those of the table style, 0.075 in).
    A caller that wants the table on a new page calls page_break first. Returns the table."""
    if title:
        p = paragraph(d, before=8, after=2)
        p.paragraph_format.keep_with_next = True             # the caption stays on the page of the table
        add_inline(p, title, bold=True)
    t = d.add_table(rows=1, cols=len(header))
    t.style = d.styles['Table Grid']
    for cell, h in zip(t.rows[0].cells, header):
        add_inline(cell.paragraphs[0], h, size=size, bold=True)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for row in rows:
        cells = t.add_row().cells
        for cell, v in zip(cells, row):
            add_inline(cell.paragraphs[0], str(v), size=size, literal=literal)
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w)
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(round(w * 1440)))
    if fixed:
        if sum(widths) > text_width(d) + 1e-9:
            raise ValueError(f'column widths sum to {sum(widths):.2f} in > {text_width(d):.2f} in (text width)')
        t.autofit = False
        tblw = t._tbl.tblPr.find(qn('w:tblW'))
        tblw.set(qn('w:type'), 'dxa')
        tblw.set(qn('w:w'), str(sum(round(w * 1440) for w in widths)))
    if cell_margin is not None:
        mar = parse_xml(f'<w:tblCellMar {nsdecls("w")}><w:left w:w="{round(cell_margin * 1440)}" w:type="dxa"/>'
                        f'<w:right w:w="{round(cell_margin * 1440)}" w:type="dxa"/></w:tblCellMar>')
        t._tbl.tblPr.insert_element_before(mar, 'w:tblLook', 'w:tblCaption', 'w:tblDescription', 'w:tblPrChange')
    if footnote:
        p = paragraph(d)
        add_inline(p, footnote, size=note_size)
        if note_spacing is not None:
            p.paragraph_format.line_spacing_rule = note_spacing
    return t


def chars_per_row(t):
    """Characters of cell text per table row (line breaks inside a cell count as one space)."""
    return [sum(len(c.text.replace('\n', ' ')) for c in r.cells) for r in t.rows]


# ------------------------------------------------------------------ figures
def figure_png(n):
    """Copy of paper/figures/Figure_<n>.png at FIGURE_DPI for FIGURE_WIDTH inches, on white, in paper/build/."""
    from PIL import Image
    src, dst = HERE / 'figures' / f'Figure_{n}.png', BUILD / f'Figure_{n}_{FIGURE_DPI}dpi.png'
    im = Image.open(src)
    if im.mode != 'RGB':
        im = im.convert('RGBA')
        white = Image.new('RGBA', im.size, (255, 255, 255, 255))
        im = Image.alpha_composite(white, im).convert('RGB')
    px = round(FIGURE_WIDTH * FIGURE_DPI)
    im.resize((px, round(im.height * px / im.width)), Image.LANCZOS).save(dst, dpi=(FIGURE_DPI, FIGURE_DPI),
                                                                          optimize=True)
    return dst


def alt_texts():
    """{'Figure 1': text, ...} from paper/figures/alt_text.txt (a heading line per image, then its text)."""
    out = {}
    for block in re.split(r'\n\s*\n', (HERE / 'figures' / 'alt_text.txt').read_text(encoding='utf-8')):
        lines = block.strip().splitlines()
        m = re.match(r'(Figure \d+|Graphical abstract) \(', lines[0]) if lines else None
        if m and len(lines) > 1:
            out[m.group(1)] = ' '.join(line.strip() for line in lines[1:])
    return out


def add_figure(d, n, legend, alt):
    """Figure n as in the template: an empty centered paragraph, the image centered and FIGURE_WIDTH wide (with its
    alternative text), and a centered 10 pt caption 'Figure n. <first sentence of the legend>'. As in the template,
    nothing keeps them together: where the image does not fit on the page, it starts the next page at the top margin
    and the empty paragraph stays behind."""
    paragraph(d, center=True)
    p = paragraph(d, center=True)
    run = p.add_run()
    run.add_picture(str(figure_png(n)), width=Inches(FIGURE_WIDTH))
    inline = run._r.find('.//' + qn('wp:inline'))
    inline.docPr.set('name', f'Figure {n}')
    if alt:
        inline.docPr.set('descr', alt)
        inline.find('.//' + qn('pic:cNvPr')).set('descr', alt)
    first = re.split(r'(?<=\.)\s+(?=[A-Z(])', plain(legend).strip(), maxsplit=1)[0]
    paragraph(d, f'Figure {n}. {first}', size=10, center=True)


# ------------------------------------------------------------------ build
def sections(text):
    """Split the body into (level, title, paragraphs) blocks by '## ' and '### ' headings."""
    blocks, cur = [], None
    for line in text.splitlines():
        if line.startswith('## ') or line.startswith('### '):
            level = 2 if line.startswith('## ') else 3
            cur = [level, line.split(' ', 1)[1].strip(), []]
            blocks.append(cur)
        elif cur is not None:
            cur[2].append(line)
    out = []
    for level, title, lines in blocks:
        paras, buf = [], []
        for line in lines + ['']:
            if line.strip():
                buf.append(line.strip())
            elif buf:
                paras.append(' '.join(buf))
                buf = []
        out.append((level, title, paras))
    return out


def words(text):
    text = re.sub(r'\^\{[^}]+\}', '', text)
    return len(re.findall(r"[A-Za-z0-9][A-Za-z0-9'’.,/–-]*", text))


def word_count_as_word(d, start='1. Introduction', stop='Study Highlights'):
    """Words of the built document from the heading `start` up to the heading `stop`, counted as Word's word count does:
    strings between spaces, with en and em dashes also separating words ('Michaelis–Menten', and the citation range
    '8–12' after 'models;'), and the section numbers ('2.1') as words. This reproduced Word's own count
    (Range.ComputeStatistics) paragraph by paragraph for the Introduction to the Conclusion of this manuscript."""
    texts = [p.text for p in d.paragraphs]
    if start not in texts or stop not in texts:
        raise SystemExit(f'word count: heading {start!r} or {stop!r} not found in the manuscript')
    i = texts.index(start)
    return sum(len([w for w in re.split(r'[\s–—]+', t) if w]) for t in texts[i:texts.index(stop, i)])


def main():
    numbers = json.loads((BUILD / 'numbers.json').read_text(encoding='utf-8'))
    table2 = json.loads((BUILD / 'table2.json').read_text(encoding='utf-8'))
    misleading = any(r['condition'].startswith('Misleading') for r in table2)
    extra = HERE / 'misleading_text.json'
    if misleading and extra.exists():
        numbers.update(json.loads(extra.read_text(encoding='utf-8')))
    numbers['ai_disclosure'] = f'⟦HL⟧{AI_DISCLOSURE}⟦/HL⟧'
    numbers['author_contributions'] = CONTRIBUTIONS
    numbers['funding'] = FUNDING                         # the Funding and Conflict of Interest Statement sections
    numbers['coi'] = COI

    text = (HERE / 'manuscript_cpt.md').read_text(encoding='utf-8')
    text = conditional_blocks(text, dict(numbers, misleading=misleading))
    text = re.sub(r'<!--.*?-->', '', text, flags=re.S)
    text = fill(text, numbers, draft='--draft' in sys.argv)
    cites = Citations()
    body = sections(text)
    body = [(lvl, title, [cites.markup(p) for p in paras]) for lvl, title, paras in body]

    main_text = [p for lvl, title, paras in _between(body, 'INTRODUCTION', 'STUDY HIGHLIGHTS') for p in paras]
    n_words = sum(words(p) for p in main_text)
    abstract = next(paras for lvl, title, paras in body if title == 'ABSTRACT')
    n_abstract = sum(words(p) for p in abstract)
    highlights = next(paras for lvl, title, paras in body if title == 'STUDY HIGHLIGHTS')
    n_highlights = sum(words(p) for p in highlights if not p.startswith('**'))
    n_refs = len(cites.order)
    n_fig, n_tab = 4, 2
    print(f'main text {n_words} words; abstract {n_abstract}; highlights {n_highlights}; references {n_refs}; '
          f'figures {n_fig}; tables {n_tab}')
    problems = []
    if n_words > 4000:
        problems.append(f'main text {n_words} > 4000 words')
    # the headings counted without the section numbers of the template ('2.1'), which CPT headings do not have; Word's
    # own count of the section in the built document (section numbers included) is printed after the build and only
    # warned about (see word_count_as_word)
    main_headings = _between(body, 'INTRODUCTION', 'STUDY HIGHLIGHTS')
    n_words_h = n_words + sum(words(title) for lvl, title, paras in main_headings)
    print(f'main text including headings {n_words_h} words (without the {len(main_headings)} section numbers)')
    if n_words_h > 4000:
        problems.append(f'main text including headings {n_words_h} > 4000 words')
    if n_abstract > 250:
        problems.append(f'abstract {n_abstract} > 250 words')
    if n_highlights >= 250:
        problems.append(f'study highlights {n_highlights} >= 250 words')
    if n_refs > 50:
        problems.append(f'{n_refs} references > 50')
    if n_fig + n_tab > 7:
        problems.append('more than 7 figures and tables')
    import check_claims
    failed = [name for name, ok, _ in check_claims.checks() if not ok]
    problems += [f'claim no longer holds: {name}' for name in failed]
    for p in problems:
        print('WARNING:', p)
    if problems and '--draft' not in sys.argv:
        raise SystemExit('final build refused (use --draft to build anyway): ' + '; '.join(problems))
    if AI_DISCLOSURE.lstrip().startswith('[To be completed'):          # reported, but the build is not refused
        print('WARNING: the AI-use disclosure is still a placeholder (AI_DISCLOSURE in paper/build_cpt.py; printed '
              'highlighted in Acknowledgments and in the cover letter); the author must write it before submission.')
    links = re.findall(r'(https://github\.com/\S+) \(tag ([^)]+)\)', text)
    print('REMINDER: before submission, these Data Availability links must exist (not checked here): '
          + '; '.join(f'{url} tag {tag}' for url, tag in links)
          + ('; the release of the first tag must carry the run logs, reference fits, and evaluation outputs.'
             if 'attached to that release' in text else '.'))

    unknown = [title for lvl, title, paras in body if lvl == 2 and title not in HEADINGS]
    if unknown:
        raise SystemExit(f'headings of paper/manuscript_cpt.md without a place in the manuscript: {unknown}')
    d = base_document()
    # title, author, affiliation, corresponding author (no separate title page, as in the template)
    paragraph(d, TITLE, size=13, bold=True, before=12, after=6)
    marks = len(AFFILIATIONS) > 1
    names = [name + (f'^{{{a}}}' if marks else '') for name, a in AUTHORS]
    paragraph(d, names[0] if len(names) == 1 else ', '.join(names[:-1]) + ', and ' + names[-1])
    for k, v in AFFILIATIONS.items():
        paragraph(d, (f'^{{{k}}}' if marks else '') + v)
    name, program, university, address, email = contact()
    paragraph(d, f'Corresponding Author: {name} ({email}), {program}, {university}, {address}')

    # abstract and keywords, then the numbered sections up to the Conclusion
    blocks = {title: (lvl, paras) for lvl, title, paras in body}

    def head(text, level):
        heading(d, text, level)
    section, n_sec, n_sub = None, 0, 0
    for lvl, title, paras in body:
        if lvl == 2:
            section = title
        if section in BACK_MATTER or section == 'FIGURE LEGENDS':
            continue
        if lvl == 2 and title in NUMBERED:
            n_sec, n_sub = n_sec + 1, 0
            head(f'{n_sec}. {HEADINGS[title]}', 2)
        elif lvl == 2:
            head(HEADINGS[title], 2)
        elif section in NUMBERED:
            n_sub += 1
            head(f'{n_sec}.{n_sub} {title_case(title)}', 3)
        else:
            head(title_case(title), 3)
        for t in paras:
            add_inline(d.add_paragraph(), t)
        if title == 'ABSTRACT':
            paragraph(d, 'Keywords: ' + '; '.join(KEYWORDS))
    for title in BACK_MATTER:
        if title not in blocks:
            continue
        head(HEADINGS[title], 2)
        for t in blocks[title][1]:
            p = d.add_paragraph()
            if t.startswith('**'):                           # a Study Highlights question stays with its answer
                p.paragraph_format.keep_with_next = True
            add_inline(p, t)
    stray = [title for lvl, title, paras in _between(body, 'STUDY HIGHLIGHTS', None) if lvl == 3]
    if stray:
        raise SystemExit(f'subsections after the Conclusion are not placed in the manuscript: {stray}')

    # references
    head('References', 2)
    for i, k in enumerate(cites.order, 1):
        add_inline(d.add_paragraph(), f'{i}. {REFERENCES[k]}')

    # tables, after a page break as in the template; the sources cited in Table 1 keep their numbers from the text
    page_break(d, empty_before=True)
    heading(d, 'Tables', 2)
    statements = json.loads((HERE.parent / 'benchmarks' / 'datasets.json').read_text(encoding='utf-8'))
    statement_note = '; '.join(f'{name}, "{statements[k]["knowledge"]}"' for k, name in
                               (('pheno', 'phenobarbital'), ('remifentanil', 'remifentanil'), ('oral_mm', 'oral MM')))
    t1 = add_table(d, '**Table 1.** Benchmark datasets, reference models, and expert statements',
                   ['Dataset', 'Design', 'Reference model'], table1_rows(cites),
                   f'Expert statements given in the expert-statement condition (verbatim): {statement_note} '
                   'CL, clearance; IV, intravenous; LBM, lean body mass; MM, Michaelis–Menten; Q2 and Q3, '
                   'intercompartmental clearances; V, V1, V2, volumes of distribution. '
                   'Details of the designs and reference models are given in Methods and Table S2.',
                   [1.3, 2.25, 2.45], fixed=True)
    if len(cites.order) != n_refs:
        raise SystemExit(f'Table 1 cites a reference not cited in the text: {cites.order[n_refs:]}')
    header = ['Dataset, condition', 'Runs', 'Reference structure', 'Reference relationships (form)',
              'Other relationships', 'Typical values agree', 'Reproduced', 'ΔOFV', 'ΔAIC']

    def two_lines(v):                                    # median (range): the range on its own line
        return str(v).replace(' (', '\n(', 1)
    rows = [[f"{r['dataset']}, {r['condition'].lower()}", r['runs'], r['structure'],
             'NA' if r['relationships'] == 'NA' else f"{r['relationships']} ({r['forms']})", two_lines(r['extra']),
             r['typical'], r['reproduced'], two_lines(r['delta_ofv']), two_lines(r['delta_aic'])]
            for r in table2]
    # 8.5 pt with the template's cell margins: the largest half-point size at which every column is at least as wide as
    # its widest word or value (Times New Roman metrics; 9 pt would need 6.1 in) within the 6.0 in; the rest of the
    # width goes where it saves the most lines
    t2 = add_table(d, '**Table 2.** Final models of the agent runs compared with the reference models',
                   header, rows,
                   'AIC, Akaike information criterion; Claude, Claude Opus 5.5; GPT, GPT-6.1 Sol; MM, '
                   'Michaelis–Menten; NA, not applicable; OFV, objective function value. '
                   'Values are counts over runs, or medians (ranges). Runs: number of runs (GPT/Claude). '
                   + ('Misleading statement: deliberately wrong statement (Supplementary Material S1); counts refer '
                      'to the true reference model. ' if misleading else '') +
                   'Reference structure: runs whose final model had the structure of the reference model. Reference '
                   'relationships (form): parameter–covariate pairs of the reference model present in the final model, '
                   'summed over runs, with the number in the reference form family in parentheses (NA, the reference '
                   'model has no covariates). Other relationships: pairs per run in the final model that are not in '
                   'the reference model. Typical values agree: the median ratio of the subject-level typical values '
                   '(final/reference model) lay within 0.80 to 1.25 for every reference parameter. Reproduced: '
                   'reference structure, exactly the reference relationships, and agreement of typical values '
                   '(stochastic models not compared). ΔOFV and ΔAIC: final model minus the PKPy2 fit of the reference '
                   'model, whose stochastic model is given in Table S2 (data scale for log-normal error).',
                   [.95, .42, .66, .80, .80, .54, .77, .53, .53], size=8.5, fixed=True)
    print('characters per row: Table 1 ' + ', '.join(map(str, chars_per_row(t1))) + '; Table 2 '
          + ', '.join(map(str, chars_per_row(t2))))
    long_rows = [f'Table {k} row {i} ({c} characters)' for k, t in ((1, t1), (2, t2))
                 for i, c in enumerate(chars_per_row(t)) if c > MAX_TABLE_ROW]
    for r in long_rows:
        print(f'WARNING: {r} > {MAX_TABLE_ROW} characters per row')
    if long_rows and '--draft' not in sys.argv:
        raise SystemExit(f'final build refused (use --draft to build anyway): table rows over {MAX_TABLE_ROW} '
                         'characters: ' + '; '.join(long_rows))

    # figure captions (the full legends), then the figures with their short captions (no graphical abstract)
    legends = blocks['FIGURE LEGENDS'][1]
    heading(d, HEADINGS['FIGURE LEGENDS'], 2)
    for t in legends:
        add_inline(d.add_paragraph(), t)
    alt = alt_texts()
    for t in legends:
        m = re.match(r'\*\*Figure (\d+)\.\*\*\s*(.*)', t)
        if not m:
            raise SystemExit(f'figure legend without "**Figure n.**": {t[:60]}')
        add_figure(d, int(m.group(1)), m.group(2), alt.get(f'Figure {m.group(1)}'))
    if len(legends) != n_fig:
        raise SystemExit(f'{len(legends)} figure legends, {n_fig} figures')
    # what an editor sees in Word: the section numbers of the template and Word's splitting at dashes add words to the
    # count checked above; reported, but the build is not refused (cutting words or keeping the numbered headings is
    # the author's decision)
    n_word_count = word_count_as_word(d)
    print(f'main text as Word counts it (1. Introduction to 5. Conclusion, with headings and section numbers): '
          f'{n_word_count} words')
    if n_word_count > 4000:
        print(f'WARNING: Word will count about {n_word_count:,} words for Introduction to Conclusion, over the CPT '
              f'limit of 4,000 (the count checked above, {n_words_h:,}, leaves out the {len(main_headings)} section '
              f'numbers and Word\'s splitting at dashes); cut about {n_word_count - 4000} words or accept the count.')

    OUT.mkdir(exist_ok=True)
    set_properties(d, TITLE)
    save_docx(d, OUT / 'PKAgent_CPT_manuscript.docx')
    cover_letter()
    for n in range(1, n_fig + 1):
        for ext in ('pdf', 'tiff'):
            src = HERE / 'figures' / f'Figure_{n}.{ext}'
            if src.exists():
                shutil.copy(src, OUT / src.name)
    shutil.copy(HERE / 'figures' / 'alt_text.txt', OUT / 'alt_text.txt')     # checked by check_claims.alt_text_checks
    (BUILD / 'manuscript_cpt_filled.md').write_text(text, encoding='utf-8')
    (BUILD / 'reference_order.json').write_text(json.dumps(cites.order, indent=1), encoding='utf-8')
    unused = sorted(set(REFERENCES) - set(cites.order))
    if unused:
        print('references not cited:', unused)
    print('wrote', OUT / 'PKAgent_CPT_manuscript.docx')


def contact():
    """Name, program, university, postal address, and email of the corresponding author."""
    rest, email = CORRESPONDING.rsplit('. Email: ', 1)
    name, program, university, address = rest.split(', ', 3)
    return name, program, university, address, email


def signature_lines():
    """Name, program, university, postal address, and email of the corresponding author, one per line."""
    name, program, university, address, email = contact()
    return [name, program, university, address, f'Email: {email}']


def cover_letter():
    """The letter in the typography of the manuscript template (Letter page, margins 1.25 in left and right and 1.0 in
    top and bottom, Times New Roman 11 pt, line spacing 1.15, 10 pt after each paragraph; no page numbers), with the
    signature block on separate lines."""
    import datetime as dt
    today = dt.date.today()
    numbers = json.loads((BUILD / 'numbers.json').read_text(encoding='utf-8'))     # run counts as in the manuscript
    values = dict(numbers, date=f'{today:%B} {today.day}, {today.year}', title=TITLE,
                  ai_disclosure=f'⟦HL⟧{AI_DISCLOSURE}⟦/HL⟧', corresponding='\n'.join(signature_lines()))
    text = (HERE / 'cover_letter_cpt.md').read_text(encoding='utf-8')
    text = re.sub(r'\{\{(\w+)\}\}', lambda m: str(values[m.group(1)]), text)
    d = base_document()
    for block in text.split('\n\n'):
        p = d.add_paragraph()
        for k, line in enumerate(block.strip().splitlines()):
            if k:
                p.add_run().add_break()
            add_inline(p, line)
    set_properties(d, 'Cover letter: ' + TITLE)
    save_docx(d, OUT / 'PKAgent_CPT_cover_letter.docx')


def _between(body, start, stop):
    out, on = [], False
    for b in body:
        if b[1] == start:
            on = True
        if b[1] == stop:
            break
        if on:
            out.append(b)
    return out


if __name__ == '__main__':
    main()
