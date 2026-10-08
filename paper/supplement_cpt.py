"""Supplementary material of the CPT submission, generated from the code, the benchmark definitions and the runs.

One file, paper/submission_cpt/Supplementary_Material.docx, with three sections:
S1  system prompt, task message, tools (Table S1), model specification schema, and the dataset descriptions, expert
    statements and misleading statements given to the agent.
S2  run and estimation settings; PKPy2 fits of the reference models (Table S2); increase in OFV when each reference
    relationship is removed (Table S3).
S3  final model, diagnostics, tests and tool use of every run (Tables S4 and S5), the deterministic covariate
    baselines, statements referring to prior knowledge (recall), and the final report of every run.
The file is written only when all three sections are built; otherwise the previous file is left unchanged.

Format: that of the author's earlier supplementary file, from paper/templates/supplement_template.docx (python
paper/templates/make_templates.py): US Letter portrait, margins 1.25 in left and right and 1.0 in top and bottom, no
header, footer or page numbers; Times New Roman 11 pt, line spacing 1.15, 10 pt after each paragraph; every paragraph in
the Normal style with direct formatting. No title page: the file starts with the heading of S1. Section headings 12 pt
bold (S2 and S3 start on a new page), subsection headings bold at body size, table captions bold, tables in the
template's grid table style with a Normal note below them: at body size where they fit the 6.0 in text width, the wide
ones at a smaller font with fixed column widths. Verbatim code (system prompt, task message, schema) is set in Consolas,
the recall excerpts and final reports at 9 pt.
Run benchmarks/reference_fits.py, benchmarks/evaluate.py, benchmarks/agent_tests.py, benchmarks/reference_table.py and
paper/manuscript_numbers.py first.
"""
import json
import re
import sys
from pathlib import Path

import pandas as pd
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.text.paragraph import Paragraph

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT / 'src'))
sys.path.insert(0, str(HERE))
from build_cpt import (TITLE, add_inline, add_table, base_document, heading, paragraph, save_docx,  # noqa: E402
                       set_properties)
from pkagent.config import Budget, Settings                           # noqa: E402
from pkagent.prompts import SYSTEM, task                              # noqa: E402
from pkagent.spec import SPEC_SCHEMA                                  # noqa: E402
from pkagent.tools import TOOLS                                       # noqa: E402
from tool_groups import N_TOOLS, TOOL_GROUPS                          # noqa: E402   (the groups of Figure 1)

OUT = HERE / 'submission_cpt'
FILE = 'Supplementary_Material.docx'
OLD_FILES = [f'Supplementary_Material_S{i}.docx' for i in (1, 2, 3)]   # the earlier layout: one file per section
SMALL = 9                                   # points: recall excerpts and final reports (verbatim)
BENCH = ROOT / 'benchmarks'
DATASETS = json.loads((BENCH / 'datasets.json').read_text(encoding='utf-8'))
LABEL = dict(pheno='Phenobarbital', remifentanil='Remifentanil', oral_mm='Oral MM (simulated)')
NAME = dict(pheno='phenobarbital', remifentanil='remifentanil', oral_mm='oral MM')       # within a sentence
LLM = dict(gpt='GPT-6.1 Sol', claude='Claude Opus 5.5')
COND = dict(none='no knowledge', knowledge='expert statement', misleading='misleading statement')
# the titles of the sections; the SUPPLEMENTARY MATERIAL list of the manuscript uses the same
TITLES = dict(
    S1='System prompt, task message, tools (Table S1), model specification schema, and the dataset descriptions, '
       'expert statements, and misleading statements given to the agent',
    S2='Run and estimation settings, PKPy2 fits of the reference models (Table S2), and the increase in OFV when each '
       'reference relationship is removed (Table S3)',
    S3='Final model, diagnostics, tests, and tool use of every run (Tables S4 and S5), the deterministic covariate '
       'baselines, statements referring to prior knowledge, and the final reports')

# abbreviations of the supplementary tables: (term, pattern, definition); a footnote lists the terms its table uses
_B, _E = r'(?<![A-Za-z0-9_])', r'(?![A-Za-z0-9_])'
ABBREVIATIONS = [(term, pattern or _B + re.escape(term) + _E, definition) for term, pattern, definition in [
    ('AGE', None, 'age'),
    ('AIC', None, 'Akaike information criterion'),
    ('APGR', _B + 'APGR', '5-minute Apgar score'),
    ('APGR<5', 'APGR<5', 'indicator for Apgar score below 5'),
    ('AUC', None, 'area under the concentration–time curve'),
    ('BIC', None, 'Bayesian information criterion'),
    ('CI', None, 'confidence interval'),
    ('CL', _B + r'CL(?![A-Za-z0-9_/])', 'clearance'),
    ('CL/F', None, 'apparent clearance'),
    ('Cmax', None, 'maximum concentration'),
    ('cmt', None, 'compartment'),
    ('CP', None, 'plasma concentration output'),
    ('CV', None, 'coefficient of variation'),
    ('CWRES', None, 'conditional weighted residuals'),
    ('DOSE', None, 'dose level'),
    ('DV', None, 'observed concentration'),
    ('FOCE-I', None, 'first-order conditional estimation with interaction'),
    ('HT', None, 'height'),
    ('IIV', None, 'interindividual variability'),
    ('IPRED', None, 'individual prediction'),
    ('IV', None, 'intravenous'),
    ('Ka', None, 'absorption rate constant'),
    ('KM', None, 'Michaelis constant'),
    ('LBM', None, 'lean body mass'),
    ('LLM', None, 'large language model'),
    ('log L', None, 'log-likelihood'),
    ('MM', None, 'Michaelis–Menten'),
    ('NCA', None, 'noncompartmental analysis'),
    ('NPDE', None, 'normalized prediction distribution errors'),
    ('OFV', None, 'objective function value'),
    ('omega^2', None, 'variance of interindividual variability'),
    ('PRED', None, 'population prediction'),
    ('Rep', None, 'replicate'),
    ('RSE', None, 'relative standard error'),
    ('SD', _B + r'SDs?' + _E, 'standard deviation'),
    ('SEX', None, 'sex'),
    ('Tmax', None, 'time of maximum concentration'),
    ('USD', None, 'US dollars'),
    ('V/F', None, 'apparent volume of distribution'),
    ('VMAX', None, 'maximum elimination rate'),
    ('VPC', None, 'visual predictive check'),
    ('WT', None, 'birth weight (phenobarbital)'),
    ('η-shr.', None, 'η-shrinkage'),
]]


def abbreviations(texts, skip=(), **override):
    """'AGE, age; CL, clearance; ...' for the terms that occur in the texts of a table (title, header, cells,
    footnote); volumes and intercompartmental clearances are grouped."""
    text = ' '.join(str(t) for t in texts)
    found = {term: override.get(term, definition) for term, pattern, definition in ABBREVIATIONS
             if term not in skip and re.search(pattern, text)}

    def present(*names):
        return [n for n in names if re.search(r'(?<![A-Za-z0-9_/])' + n + r'(?![A-Za-z0-9_/])', text)]
    vols = present('V', 'V1', 'V2', 'V3')
    if vols:
        numbered = [v for v in vols if v != 'V']
        names = (['V'] if 'V' in vols else []) + (['V1–V3'] if len(numbered) == 3 else numbered)
        found[', '.join(names)] = 'volume of distribution' if len(vols) == 1 else 'volumes of distribution'
    qs = present('Q', 'Q2', 'Q3')
    if qs:
        found[' and '.join(qs) if len(qs) == 2 else ', '.join(qs)] = ('intercompartmental clearance' if len(qs) == 1
                                                                       else 'intercompartmental clearances')
    order = sorted(found, key=lambda t: (not t[0].isascii(), t.lower()))
    return '; '.join(f'{t}, {found[t]}' for t in order) + '.'


def run_name(ds, cond, llm, rep):
    return f'{NAME[ds]}, {COND[cond]}, {LLM[llm]}, replicate {rep[3:]}'


def run_dir(ds, cond, llm, rep):
    return BENCH / 'runs' / ds / cond / llm / rep


def tool_log(path):
    return [json.loads(line) for line in (path / 'tool_log.jsonl').read_text(encoding='utf-8').splitlines()
            if line.strip()]


def text(d, t, small=False):
    """A Normal paragraph with the markup of add_inline: the template's 11 pt, line spacing 1.15 and 10 pt after, or
    9 pt with 4 pt after (small)."""
    return add_inline(paragraph(d, after=4 if small else None), t, size=SMALL if small else None)


def verbatim(d, t):
    """Verbatim text (no markup) at 9 pt with 4 pt after."""
    return add_inline(paragraph(d, after=4), t, size=SMALL, literal=True)


def labeled(d, label, body, small=False, bold=False):
    """A paragraph of an italic (or bold) label followed by verbatim text (no markup: '*' and '**' are kept), at body
    size or, small, at 9 pt with 4 pt after."""
    size = SMALL if small else None
    p = paragraph(d, after=4 if small else None)
    run = p.add_run(label)
    if bold:
        run.bold = True
    else:
        run.italic = True
    if size:
        run.font.size = Pt(size)
    add_inline(p, body, size=size, literal=True)
    return p


def group_label(d, t):
    """Bold label of a group of paragraphs (a dataset, a run): body size, 8 pt before and 2 pt after, kept with the
    paragraph that follows."""
    return paragraph(d, t, bold=True, before=8, after=2, keep_next=True)


def code(d, t, size=8):
    """Verbatim code: one single-spaced Consolas paragraph per line, without space between the lines and with the
    template's 10 pt after the block."""
    p = None
    for line in t.rstrip().splitlines():
        p = d.add_paragraph()
        run = p.add_run(line if line else ' ')
        run.font.name = 'Consolas'
        run.font.size = Pt(size)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(0)
    if p is not None:
        p.paragraph_format.space_after = Pt(10)


def section(d, key):
    """Heading of a section, 12 pt bold as the headings of the template; S2 and S3 start on a new page."""
    heading(d, f'Supplementary Material {key}. {TITLES[key]}', 2, new_page=key != 'S1')


def table(d, title, header, rows, footnote, widths, **kw):
    """Caption, table and note by add_table of build_cpt (the template's grid table style, header row repeated on each
    page), with fixed column widths, which must fit the 6.0 in text width; the caption is kept on the page of the
    table, and a row is not split across pages."""
    t = add_table(d, title, header, rows, footnote, widths, fixed=True, **kw)
    Paragraph(t._tbl.getprevious(), d._body).paragraph_format.keep_with_next = True
    for row in t.rows:
        row._tr.get_or_add_trPr().insert(0, OxmlElement('w:cantSplit'))
    return t


def s1(d):
    section(d, 'S1')
    heading(d, 'S1.1 System prompt', 3)
    text(d, 'Given verbatim to the language model at the start of every run.')
    code(d, SYSTEM)
    heading(d, 'S1.2 Task message', 3)
    b = Budget()
    text(d, 'The first user message of a run. The section "Expert knowledge from the analyst" appears only in the '
            'expert-statement and misleading-statement conditions; the Budget section gives the limits of the benchmark '
            f'on fits, responses, and hours, and the session also stops a run at ${b.max_cost_usd:g} of language model '
            'fees, which the message does not state (S2.1).')
    code(d, task('<dataset description>', '<expert statement>', None, b))
    heading(d, 'S1.3 Tools', 3)
    definitions = {t['function']['name']: t['function'] for t in TOOLS}
    assert N_TOOLS == len(TOOLS) == len(definitions), 'paper/tool_groups.py does not cover the tools of src/pkagent'
    rows, spans = [], []                 # spans: first and last body row of each group (row 0 is the header)
    for group, names in TOOL_GROUPS.items():
        spans.append((len(rows) + 1, len(rows) + len(names)))
        for i, name in enumerate(names):
            f = definitions[name]
            params = ', '.join(f['parameters'].get('properties', {}))
            rows.append([group if i == 0 else '', name, f['description'], params or '–'])
    title = (f'**Table S1.** The {N_TOOLS} tools available to the language model, by functional group '
             '(function-calling definitions sent with every request; descriptions verbatim)')
    header = ['Group', 'Tool', 'Description', 'Parameters']
    note = ('Group, the functional group of the tool in Figure 1 (paper/tool_groups.py); the grouping is descriptive '
            'and is not sent to the model. The model specification accepted by fit_models is given in S1.4.')
    tbl = table(d, title, header, rows,
                f'{note} {abbreviations([title, *header, *(c for r in rows for c in r)], WT="weight")}',
                [0.78, 1.27, 2.68, 1.27], size=9, literal=True)
    for first, last in spans:            # one group label per group: merge its cells of the Group column
        if last > first:
            tbl.cell(first, 0).merge(tbl.cell(last, 0))
    heading(d, 'S1.4 Model specification schema', 3)
    text(d, 'JSON schema of one model specification (an item of fit_models.models). Specifications are validated '
            'against this schema and additional semantic checks before fitting.')
    code(d, json.dumps(SPEC_SCHEMA, indent=1), size=7)
    heading(d, 'S1.5 Benchmark datasets: descriptions, expert statements, and misleading statements given to the '
               'agent', 3)
    for name in LABEL:
        v = DATASETS[name]
        group_label(d, LABEL[name])
        labeled(d, 'Description given to the agent: ', v['description'])
        labeled(d, 'Expert statement (expert-statement condition): ', v['knowledge'])
        if v.get('misleading_knowledge'):
            labeled(d, 'Misleading statement (misleading-statement condition): ', v['misleading_knowledge'])


def over_time(max_hours):
    """Clause on the runs that ended after the hour limit (the limits are checked between steps)."""
    runs = pd.read_csv(BENCH / 'evaluation' / 'runs.csv')
    over = runs[runs['hours'] > max_hours].sort_values('hours')
    clause = ('; the limits are checked before each LLM response and each fit, and a response or fit in progress is '
              'completed')
    if over.empty:
        return clause
    minutes = [round((h - max_hours) * 60) for h in over['hours']]
    names = '; '.join(run_name(r['dataset'], r['condition'], r['llm'], r['rep']) for _, r in over.iterrows())
    if len(over) == 1:
        return (clause + f', so one run ({names}) ended {minutes[0]} minute{"s" if minutes[0] != 1 else ""} after the '
                f'{max_hours:g}-hour limit')
    return clause + f', so {len(over)} runs ({names}) ended up to {max(minutes)} minutes after the {max_hours:g}-hour limit'


def reference_fit(name):
    path = BENCH / 'reference_fits' / name / 'reference_fit.json'
    if not path.exists():
        raise SystemExit(f'{path.relative_to(ROOT).as_posix()} does not exist (S2 needs the oral MM reference fit with '
                         'log-normal error in reference_fits/oral_mm and the earlier fit with proportional error in '
                         'reference_fits/oral_mm_proportional): finish "python benchmarks/reference_fits.py oral_mm", '
                         'then rerun benchmarks/reference_table.py, evaluate.py and agent_tests.py')
    return json.loads(path.read_text(encoding='utf-8'))


def oral_mm_residual():
    """S2 item on the residual model of the oral MM reference fit; the OFV difference is read from the two fits."""
    new, old = reference_fit('oral_mm'), reference_fit('oral_mm_proportional')
    for fit, kind in ((new, 'lognormal'), (old, 'proportional')):
        if fit.get('status') != 'converged' or f'CP:{kind}' not in (fit.get('description') or ''):
            raise SystemExit(f'oral MM reference fit with {kind} error not as expected: {fit.get("description")} '
                             f'({fit.get("status")})')
    delta = old['ofv'] - new['ofv']
    if delta <= 0:
        raise SystemExit(f'the proportional-error oral MM reference fit has the lower OFV ({old["ofv"]} vs '
                         f'{new["ofv"]}): revise the S2 text on the oral MM residual model')
    chk = json.loads((BENCH / 'evaluation' / 'oral_mm_residual_check.json').read_text(encoding='utf-8'))
    runs = pd.read_csv(BENCH / 'evaluation' / 'runs.csv')
    om = runs[(runs['dataset'] == 'oral_mm') & runs['condition'].isin(['none', 'knowledge'])]
    n_prop = int((om['delta_ofv_vs_reference'].round(3) == round(delta, 3)).sum())
    n_logn = int((om['delta_ofv_vs_reference'].abs() < 1e-3).sum())
    if n_prop + n_logn != len(om):
        raise SystemExit('oral MM final models are neither the proportional nor the log-normal reference model: '
                         'revise the S2 text on the oral MM residual model')
    log_skew = f"{chk['log_skewness']:.2f}".replace('-', '−')
    return ('Oral MM residual model', 'The source documentation gives 20% residual error without its type. Relative to '
            'the true individual predictions computed from the individual parameters in the source file (first '
            f"{chk['subjects']} subjects, {chk['observations']:,} observations; benchmarks/residual_check.py), "
            f"log(DV/IPRED) had SD {chk['log_sd']:.3f} and skewness {log_skew}, whereas DV/IPRED − 1 "
            f"had skewness {chk['ratio_skewness']:.2f} ({chk['expected_ratio_skewness_exponential']:.2f} expected for "
            '20% exponential error); the reference model therefore uses log-normal error. This residual model was '
            'identified after all runs had finished; the benchmark definition frozen before the runs (commit f5a4263) '
            f'used proportional error, under which the {n_prop} final models with proportional error (runs without '
            f'knowledge or with the expert statement) matched the reference OFV and the {n_logn} with log-normal error '
            f'were {delta:,.1f} lower. With proportional error, the OFV of the PKPy2 reference fit was {delta:,.1f} '
            'higher; that fit is kept in benchmarks/reference_fits/oral_mm_proportional')


def numbers():
    return json.loads((HERE / 'build' / 'numbers.json').read_text(encoding='utf-8'))


def robustness():
    """S2 item on single-start fits (numbers from paper/manuscript_numbers.py)."""
    n = numbers()
    return ('Robustness of single-start fits', 'converged fits of all runs were grouped by dataset and by specification '
            'apart from the starting values (structure, interindividual variability and covariance blocks, residual '
            'model, covariate relationships with form, center, bounds and fixed coefficients, fixed parameters); fits '
            f"with identical starting values had identical OFVs; of {n['start_groups']} groups with different starting "
            f"values, {n['start_spread_over1']} differed by more than 1 OFV unit, at most by {n['start_spread_max']} "
            f"(next {n['start_spread_next']}); within each run, a converged model with one more residual-error component "
            'or interindividual variability term than another, otherwise identical, model had a higher OFV in '
            f"{n['nested_inversions']} case (by {n['nested_inversion_max']})")


def clinical():
    """S2 item on the clinical calculations of the Discussion (numbers from paper/manuscript_numbers.py)."""
    n = numbers()
    return ('Clinical calculations', 'phenobarbital: typical concentration after a 20 mg/kg loading dose = dose / '
            f"typical V, with V per kg and the Apgar effect of the reference ({n['apgar_conc_low']} mg/L for an Apgar "
            'score below 5), divided for each run without knowledge by its median ratio of typical V to the reference '
            'in the subjects with an Apgar score below 5; remifentanil: plasma 50% and 80% decrement times after 1-, 4- '
            'and 10-hour constant-rate infusions in the published model (LBM 55 kg; ages 65, 75 and 85 years) with V3 '
            f"as published and divided by 4 (Q3 {n['v3_q3_cl_pct']}% of CL; decrement times differ by at most "
            f"{n['v3_decrement_diff_max']} minutes and end-of-infusion concentrations by at most "
            f"{n['v3_conc_diff_max']}%)")


def s2(d):
    section(d, 'S2')
    st, b = Settings(), Budget()
    heading(d, 'S2.1 Settings of the benchmark runs', 3)
    items = [
        ('Language models', 'openai/gpt-6.1-sol and anthropic/claude-opus-5.5 through OpenRouter; reasoning effort '
                            f'"{st.reasoning_effort}"; provider-default temperature and routing (not seeded, not '
                            f'recorded); at most {st.max_output_tokens:,} output tokens per response; explicit prompt '
                            'caching for Claude Opus 5.5'),
        ('Plot images', 'the plots of one round of tool calls (at most six) form one message; the images of the two most '
                        'recent such messages are kept, older ones are replaced by a note'),
        ('Budgets per run', f'{b.max_fits} fits (bootstrap refits excluded), {b.max_turns} LLM responses, '
                            f'{b.max_hours:g} hours, {b.max_cost_usd:g} USD of LLM fees; the task message states the '
                            'first three, and every tool result reports the remaining fits, responses and hours and the '
                            'fees spent' + over_time(b.max_hours)),
        ('Computing', 'one desktop computer (AMD Ryzen 5 5600, 6 cores, 12 threads, 32 GB memory, Windows 11); two '
                      'runs at a time, each with three fitting processes of two Numba threads; other computations ran '
                      'on the same computer'),
        ('Estimation', "PKPy2 0.2.1, method='laplace': Laplace approximation of the marginal likelihood with "
                       'expected-information (FOCE-I-type) individual curvature'),
        ('Optimizer', 'one start (the agent\'s values); L-BFGS-B, at most 150 iterations (ftol 1e-10, gtol 1e-4); '
                      'covariate coefficients scaled by the standard deviation of their covariate; linear covariate '
                      'coefficients bounded so that 1 + β(z − center) stays positive for every subject (margin 0.98)'),
        ('Gradient', 'central finite differences of the Laplace objective (relative step 1e-4) with the conditional modes '
                     're-solved at every difference point; warm starts anchored at the best point'),
        ('Conditional modes', 'Newton search; a mode solved without a warm start receives a second search from the best '
                              'of 256 scrambled Sobol draws from N(0, Ω), and the mode with the higher joint density is '
                              'kept; at the end of a minimization, cold and warm modes are compared and the minimization '
                              'continues (at most twice) when the cold modes lower the OFV; modes are saved with the fit'),
        ('Time limit', f'{st.fit_wall_seconds / 60:g} minutes for the L-BFGS-B minimization; the polish (at most 20 '
                       'iterations), convergence check, standard errors and diagnostics are not time-limited'),
        ('Convergence', f'Newton decrement of the OFV at most {st.laplace_tolerance}'),
        ('Standard errors', 'covariance 2H⁻¹, with H from central differences of the gradient (modes re-solved) at '
                            'relative steps 0.002 and 0.001; reported when H is positive definite, the two Hessians '
                            'differ by at most 10% (relative norm), the standard errors agree within 20%, each Hessian '
                            'is at most 10% asymmetric, and no step crosses a bound'),
        ('Bootstrap (on request)', f'parallel Laplace refits, at most 200 samples, time limit '
                                   f'{st.bootstrap_seconds / 60:g} minutes'),
        ('Log-normal residual models', 'PKAgent reports the OFV on the data scale (PKPy2 OFV plus twice the sum of the '
                                       'log observations), the −2 log-likelihood of the observations themselves, so it '
                                       'is comparable with additive, proportional, and combined error models; the '
                                       'agents saw these values'),
        ('Reference fits', 'same session code without an LLM; two starts (the published values and one perturbation); '
                           '60-minute limit'),
        ('Reference-fit VPCs', 'the VPC of the final models (500 simulations, 8 bins), without and with prediction '
                               f'correction, with the default seed {st.seed} (the runs used their own seeds; '
                               'benchmarks/reference_vpc.py)'),
        oral_mm_residual(),
        ('Removal of reference relationships', 'one start from the reference estimates; 60-minute limit; no standard '
                                               'errors; the phenobarbital weight exponents are fixed, so their OFV '
                                               'changes are descriptive'),
        ('Stepwise baseline', 'phenobarbital: reference structure without covariates; forward inclusion at P < 0.05 and '
                              'backward elimination at P < 0.01 over weight on CL and V (power function, exponent '
                              'estimated) and an indicator for Apgar score below 5 on CL and V; remifentanil: backward '
                              'elimination at P < 0.01 (OFV increase 6.63) from the PKPy2 fit of the reference model, '
                              'followed by removal of each remaining relationship in turn'),
        ('Development', '16 pilot runs (10 completed) on the phenobarbital and remifentanil datasets under earlier tool '
                        'versions were discarded; the system prompt was not changed after the first pilot; the code '
                        'was frozen at PKAgent commit f5a4263 and PKPy2 0.2.1 (commit 41301f1 of '
                        'https://github.com/Gumgo91/PKPy2) before the benchmark; covariate recall, precision and form '
                        'agreement and the misleading-statement condition were added to the evaluation after the first '
                        'replicate, and the oral MM reference residual model was changed from proportional to '
                        'log-normal after all runs (see Oral MM residual model)'),
        robustness(),
        clinical(),
    ]
    for k, v in items:
        labeled(d, f'{k}: ', f'{v}.', bold=True)
    heading(d, 'S2.2 PKPy2 fits of the reference models', 3)
    ref = pd.read_csv(BENCH / 'evaluation' / 'reference_table.csv')
    for name in LABEL:                          # the table must come from the current reference fits
        ofv = ref[(ref['dataset'] == name) & ref['quantity'].str.startswith('OFV')]['pkpy2']
        fit = reference_fit(name)
        if ofv.empty or abs(float(ofv.iloc[0]) - fit['ofv']) > 1e-3:
            raise SystemExit(f'benchmarks/evaluation/reference_table.csv is stale for {name} (OFV '
                             f'{None if ofv.empty else float(ofv.iloc[0])} vs {fit["ofv"]} in the reference fit): '
                             'rerun benchmarks/reference_table.py')
    rows = []
    for _, r in ref.iterrows():
        def f(x, digits=4):
            return '–' if pd.isna(x) else f'{x:.{digits}g}'
        rows.append([LABEL.get(r['dataset'], r['dataset']), r['quantity'].replace('ug/', 'µg/'), f(r['reference']),
                     f(r['pkpy2'], 5), f(r['ratio'], 3), f(r['rse_percent'], 3)])
    title = ('**Table S2.** Reference values (phenobarbital: FOCE-I estimates of the NONMEM example model from a '
             'NONMEM 7.4.2 run distributed with the Pharmpy test data, pheno_real.mod and pheno_real.ext, '
             'https://github.com/pharmpy/pharmpy; remifentanil: Minto et al.; oral MM: nominal simulation values) '
             'and the PKPy2 fit of the same model to the benchmark data')
    header = ['Dataset', 'Quantity', 'Reference', 'PKPy2', 'Ratio', 'RSE (%)']
    note = ('The NONMEM objective of the phenobarbital model (586.276) differs from the PKPy2 OFV by the constant '
            '155·ln(2π). Remifentanil slopes are absolute changes per year of age or per kg of LBM; the additive '
            'published model was fitted as a product of linear terms, and its variability model (exponential IIV on '
            'all six parameters, proportional error) is not from the publication; its fit had two optima 0.6 apart '
            'and no standard errors; its values are at age 40 years and LBM 55 kg, whereas the text gives medians over '
            'subjects. Oral MM: exponential IIV on Ka, V, VMAX, and KM, as in the simulation, and exponential '
            '(log-normal) residual error (S2.1, oral MM residual model); geometric means of the simulated individual '
            'values in the subset were Ka 0.98 1/h, V 67.2 L, VMAX 981 µg/h, KM 232 µg/L.')
    table(d, title, header, rows,
          note + ' ' + abbreviations([title, *header, *(c for r in rows for c in r), note],
                                     RSE='relative standard error (– when not available)'),
          [1.05, 1.75, .9, .95, .6, .75])
    ev_path = BENCH / 'evaluation' / 'effect_evidence.json'
    if ev_path.exists():
        ev = json.loads(ev_path.read_text(encoding='utf-8'))
        rows = [[LABEL.get(ds, ds), r['effect'].replace('APGR_LT5', 'APGR<5'), r['status'],
                 '–' if r['delta_ofv'] is None else f"{r['delta_ofv']:.2f}"]
                for ds, v in ev.items() for r in v['removals']]
        heading(d, 'S2.3 Likelihood evidence of the reference covariate relationships', 3)
        title = ('**Table S3.** Increase in OFV when one covariate relationship is removed from the reference model '
                 'and the model is refitted')
        header = ['Dataset', 'Relationship removed', 'Fit status', 'ΔOFV']
        note = ('Thresholds of the agent: 3.84 for inclusion (P < 0.05) and 6.63 for retention (P < 0.01). The '
                'phenobarbital weight exponents are fixed at 1, so their OFV changes are descriptive. Relationships '
                'are written parameter~covariate.')
        table(d, title, header, rows,
              note + ' ' + abbreviations([title, *header, *(c for r in rows for c in r), note]),
              [1.5, 2.1, 1.3, 1.1])


APGR_LT5 = re.compile(r'^\s*(where\(\s*)?APGR\s*<\s*5(?![0-9.])')     # agent-made indicators of Apgar score < 5
REPORT_FIELDS = ('summary', 'development', 'expert_knowledge', 'evaluation', 'limitations')


def derived_columns(log):
    """{column: expression} of the columns the agent added to the data."""
    return {x['args'].get('name'): x['args'].get('expression') or '' for x in log if x['tool'] == 'add_data_column'}


def parameter_first(description, derived):
    """Relationships of a runs.csv description as parameter~covariate(form) ('WT~CL(power)' -> 'CL~WT(power)');
    an agent-made indicator of Apgar score below 5 is written APGR<5 ('APGRLOW=1.0~V(categorical)' -> 'V~APGR<5')."""
    def repl(m):
        cov, level, par, form = m.groups()
        if form == 'categorical' and APGR_LT5.match(derived.get(cov, '')) and (level or '=1') in ('=1', '=1.0'):
            return f'{par}~APGR<5'
        return f'{par}~{cov}{level or ""}({form})'
    return re.sub(r'(?<![A-Za-z0-9_~])([A-Za-z][A-Za-z0-9_]*)(=[^~,;()\s]+)?~([A-Za-z][A-Za-z0-9]*)\(([a-z_]+)\)',
                  repl, description)


def test_label(x):
    """Label of a covariate test in Table S5: relationship(form); the Apgar indicator as APGR<5."""
    if x['form'] == 'categorical' and x['relationship'].endswith('~APGR'):
        return x['relationship'] + '<5'
    return f"{x['relationship']}({x['form']})"


def own_final_vpc(log, final_id):
    """Bins of the agent's last VPC of the final model if it was not prediction-corrected (PKAgent then used it as the
    final VPC); None if PKAgent ran the final VPC after finalization."""
    last = None
    for x in log:
        if (x['tool'] == 'run_vpc' and x['args'].get('model_id') == final_id and isinstance(x.get('result'), dict)
                and 'error' not in x['result'] and x['result'].get('status') != 'error'):
            last = x['args']
    if last is None or last.get('prediction_corrected'):
        return None
    return int(last.get('bins') or 8)


def final_fit_without_se(log, final_id):
    """Whether the fit_models call that produced the final model had standard_errors=false."""
    for x in log:
        if x['tool'] == 'fit_models' and isinstance(x.get('result'), dict):
            if final_id in [m.get('model_id') for m in x['result'].get('results') or []]:
                return x['args'].get('standard_errors', True) is False
    return False


def clean_context(context, term):
    """Recall snippet without raw tool-call JSON: cut after the match at the first tool call ('[{"id": "toolu_...'),
    drop a tool-call list that ends before the match, and remove call ids when the match is inside the arguments."""
    starts = [m.start() for m in re.finditer(re.escape(term), context)]
    pos = min(starts, key=lambda s: abs(s - 120)) if starts else 0
    head, tail = context[:pos], context[pos:]
    m = re.search(r',?\s*\[?\{"id": "', tail)
    if m:
        tail = tail[:m.start()].rstrip()
    end = head.rfind('}}]')
    if end >= 0:
        head = head[end + 3:].lstrip()
    head = re.sub(r'^[^"\s]*", "type": "function", ', '', head)          # tail of a call id cut by the context
    head = re.sub(r'\{"id": "[^"]*", "type": "function", ', '{', head)
    return head + tail


def check_reference_ofv(runs):
    """ΔOFV in runs.csv must be relative to the current reference fits (evaluate.py rerun after a refit)."""
    for name, g in runs.groupby('dataset'):
        implied = (g['ofv'] - g['delta_ofv_vs_reference']).dropna()
        path = BENCH / 'reference_fits' / name / 'reference_fit.json'
        if not path.exists():
            print(f'WARNING: {path.relative_to(ROOT).as_posix()} missing; the OFV differences of {name} in Table S4 '
                  'were not checked against it')
            continue
        ofv = json.loads(path.read_text(encoding='utf-8'))['ofv']
        if implied.empty or (implied - ofv).abs().max() > 1e-3:
            raise SystemExit(f'benchmarks/evaluation/runs.csv is stale for {name}: its OFV differences refer to a '
                             f'reference OFV of {implied.round(3).unique().tolist()}, the reference fit has {ofv}; '
                             'rerun benchmarks/evaluate.py')


def s3(d):
    section(d, 'S3')
    st, b = Settings(), Budget()
    runs = pd.read_csv(BENCH / 'evaluation' / 'runs.csv')
    order = dict(dataset=['pheno', 'remifentanil', 'oral_mm'], condition=['none', 'knowledge', 'misleading'],
                 llm=['gpt', 'claude'])                     # the order of Table 2
    runs = runs.sort_values(['dataset', 'condition', 'llm', 'rep'],
                            key=lambda c: c.map(order[c.name].index) if c.name in order else c)
    check_reference_ofv(runs)
    keys = [(r['dataset'], r['condition'], r['llm'], r['rep']) for _, r in runs.iterrows()]
    tests = {(t['dataset'], t['condition'], t['llm'], t['rep']): t
             for t in json.loads((BENCH / 'evaluation' / 'agent_tests.json').read_text(encoding='utf-8'))}
    missing = [k for k in keys if k not in tests]
    assert not missing, f'runs without an entry in benchmarks/evaluation/agent_tests.json: {missing}'
    logs = {k: tool_log(run_dir(*k)) for k in keys}
    sv = BENCH / 'evaluation' / 'standard_vpc.json'
    standard_vpc = json.loads(sv.read_text(encoding='utf-8')) if sv.exists() else {}

    def f(x, digits=1):
        return '–' if x is None or (isinstance(x, float) and pd.isna(x)) else f'{x:,.{digits}f}'.replace('-', '−')

    # Table S4: final model, diagnostics, resources; the run profile (SEs, bootstrap, hours left) from agent_tests
    rows, own_vpc, se_after = [], [], []
    for key, (_, r) in zip(keys, runs.iterrows()):
        t, log = tests[key], logs[key]
        bins = own_final_vpc(log, r['model_id'])
        res = json.loads((run_dir(*key) / 'results.json').read_text(encoding='utf-8'))
        vpc = ((res.get('final_model') or {}).get('vpc') or {})
        used = [len(v['bins']) for v in (vpc.get('result') or {}).values() if isinstance(v, dict) and 'bins' in v]
        if vpc.get('prediction_corrected') or set(used) != {bins or 8}:
            raise SystemExit(f'final VPC of {key} not as described in the Table S4 footnote: {used} bins, agent VPC '
                             f'{bins}')
        if bins is not None:
            if '/'.join(key) not in standard_vpc:
                raise SystemExit(f"{key}: the agent's own VPC was kept; run benchmarks/standard_vpc.py and "
                                 'benchmarks/evaluate.py')
            own_vpc.append((key, bins))
        if final_fit_without_se(log, r['model_id']):
            se_after.append((key, t['final_uncertainty']))
        rows.append([LABEL[r['dataset']], COND[r['condition']], LLM[r['llm']], r['rep'][3:],
                     parameter_first(r['description'], derived_columns(log)),
                     f(r['delta_ofv_vs_reference']), f(r['delta_aic_vs_reference']),
                     f(r['vpc_inside_fraction'] * 100 if pd.notna(r['vpc_inside_fraction']) else None, 0),
                     f(r['max_rse_structural']), f(r['max_eta_shrinkage'], 0),
                     'yes' if t['final_uncertainty'] == 'computed' else 'no',
                     ', '.join(f"{x['converged']}/{x['requested']}" for x in t['bootstraps']) or '–',
                     int(r['fits']), int(r['llm_calls']), f(r['hours'], 2), f(t['hours_left'], 2),
                     f(r['cost_usd'], 2)])

    def which(items):
        return ('one run' if len(items) == 1 else f'{len(items)} runs') + f" ({'; '.join(run_name(*k) for k in items)})"
    vpc_note = ('VPC coverage is from the VPC without prediction correction that PKAgent ran after finalization (500 '
                'simulations, 8 bins, the seed of the run)')
    if own_vpc:
        vpc_note += (f"; in {which([k for k, _ in own_vpc])}, PKAgent had kept the agent's own last VPC of the final "
                     f"model ({', '.join(sorted({f'{n} bins' for _, n in own_vpc}))}), and the same standard VPC was run "
                     'afterwards (benchmarks/standard_vpc.py)')
    se_note = ''
    if se_after:
        done = all(u == 'computed' for _, u in se_after)
        se_note = (f' In {which([k for k, _ in se_after])}, the final model had been fitted without standard errors '
                   f"and PKAgent {'computed them' if done else 'attempted to compute them'} after finalization.")
    title = '**Table S4.** Final model, diagnostics, and resources of every run'
    header = ['Dataset', 'Condition', 'LLM', 'Rep', 'Final model', 'ΔOFV', 'ΔAIC', 'VPC (%)', 'Max RSE (%)',
              'Max η-shr. (%)', 'SEs', 'Bootstrap', 'Fits', 'LLM resp.', 'Hours', 'Hours left', 'USD']
    note = ('ΔOFV and ΔAIC, final model minus the PKPy2 fit of the reference model; VPC (%), percentage of observed '
            '5th, 50th and 95th percentiles inside their simulated 95% intervals (not prediction-corrected); Max RSE, '
            'largest relative standard error of a typical value; Max η-shr., largest η-shrinkage; SEs, standard '
            'errors of the final model available; Bootstrap, converged/requested replicates within the '
            f'{st.bootstrap_seconds / 60:g}-minute tool limit; Hours left, of the {b.max_hours:g}-hour budget at '
            'finalization; CP, plasma concentration output; Rep, replicate; relationships are written '
            f'parameter~covariate(form). {vpc_note}.{se_note}')
    # 17 columns on the portrait page: 6 pt with narrow cell margins; each column is as wide as its longest word (the
    # header words included), and the final model takes the rest of the 6.0 in
    table(d, title, header, rows,
          note + ' ' + abbreviations([title, *header, *(c for r in rows for c in r)],
                                     skip=('CP', 'Rep', 'RSE', 'η-shr.')),
          [.54, .44, .37, .22, 1.06, .31, .31, .25, .26, .29, .21, .43, .21, .26, .30, .30, .24], size=6,
          literal=True, cell_margin=.025)

    # Table S5: tool use and covariate tests
    rows = []
    for key in keys:
        t, derived = tests[key], derived_columns(logs[key])
        if any(x['form'] == 'categorical' and x['relationship'].endswith('~APGR') for x in t['covariate_tests']):
            bad = {c: e for c, e in derived.items() if 'APGR' in e and not APGR_LT5.match(e)}
            if bad:
                raise SystemExit(f'{key}: Apgar column other than an indicator of APGR < 5: {bad}')
        ev = {}
        for x in t['covariate_tests']:
            ev.setdefault(test_label(x), []).append(x['delta_ofv'])
        ev_text = '; '.join(f"{k} {', '.join(f(v, 1) for v in vs)}" for k, vs in sorted(ev.items())) or '–'
        rows.append([LABEL[key[0]], COND[key[1]], LLM[key[2]], key[3][3:],
                     'yes' if t['plotted_data'] else 'no', 'yes' if t['ran_nca'] else 'no',
                     'yes' if t['screened_covariates'] else 'no', 'yes' if t['fitted_after_viewing_plots'] else 'no',
                     ev_text])
    title = '**Table S5.** Tool use and the covariate tests of every run'
    header = ['Dataset', 'Condition', 'LLM', 'Rep', 'Plotted data', 'NCA', 'Screened covariates',
              'Fitted after viewing plots', 'Covariate tests (ΔOFV)']
    note = ('Plotted data, NCA, Screened covariates: the tool was called at least once; Fitted after viewing plots: a '
            'fit_models or covariate_search call after the first view_plots call. Covariate tests: OFV differences '
            'between two converged models of the run that differ only by the relationship (positive values favor the '
            'relationship), extracted from the model registry; relationships are written parameter~covariate, form in '
            'parentheses; APGR<5, indicator for Apgar score below 5; for oral MM, CL~DOSE was tested in models with '
            'linear elimination, the other DOSE relationships in Michaelis–Menten models.')
    table(d, title, header, rows,
          note + ' ' + abbreviations([title, *header, *(c for r in rows for c in r)], skip=('APGR<5',)),
          [.8, .67, .57, .37, .52, .42, .66, .55, 1.44], size=8, literal=True)

    scm_path = BENCH / 'evaluation' / 'scm_baseline.json'
    bb_path = BENCH / 'evaluation' / 'backward_baseline.json'
    if scm_path.exists() or bb_path.exists():
        heading(d, 'S3.1 Deterministic covariate baselines', 3)
    if scm_path.exists():
        for ds, v in json.loads(scm_path.read_text(encoding='utf-8')).items():
            steps = [h for h in v['history'] if h['step'] in ('forward add', 'forward stop', 'backward remove')]
            desc = '; '.join(f"{h['step']} {h.get('added') or h.get('removed') or h.get('best')} (ΔOFV "
                             f"{h['delta_ofv']})" for h in steps)
            labeled(d, f'{LABEL[ds]} (stepwise covariate modeling). ',
                    (f"Base OFV {v['base_ofv']:.2f}; final OFV {v['final_ofv']:.2f}; retained: "
                     f"{', '.join(v['included']) or 'none'}; {v['fits']} fits. Steps: {desc}.")
                    .replace('APGR_LT5=1(categorical)', 'APGR<5'), bold=True)
    if bb_path.exists():
        for ds, v in json.loads(bb_path.read_text(encoding='utf-8')).items():
            steps = '; '.join(f"step {h['step']}: " + ', '.join(f"{k} {x:.1f}" for k, x in h['tests'].items() if x is not None)
                              for h in v['history'])
            labeled(d, f'{LABEL[ds]} (backward elimination from the reference model). ',
                    f"Removed: {', '.join(v.get('removed', [])) or 'none'}; retained: "
                    f"{', '.join(v.get('retained', []))}. OFV increase on removal at each step: {steps}.", bold=True)

    heading(d, 'S3.2 Statements referring to prior knowledge of the data or their analysis', 3)
    recall = json.loads((HERE / 'build' / 'recall.json').read_text(encoding='utf-8'))
    published = runs[runs['dataset'].isin(['pheno', 'remifentanil'])]
    main_runs = published[published['condition'].isin(['none', 'knowledge'])]
    assert set(recall) == {'/'.join(k) for k in map(tuple, main_runs[['dataset', 'condition', 'llm', 'rep']].values)}, \
        'paper/build/recall.json does not cover the runs without knowledge and with the expert statement'
    other = published[(published['condition'] == 'misleading') & published['recall_terms'].notna()]
    scope = (f'Matches of the regular expression \\bclassic\\b|well[- ]known|textbook|nonmem example|published|literature|'
             f'grasela|donn\\b|minto (case-insensitive) in the assistant messages, tool-call arguments and reasoning '
             f'summaries of the {len(recall)} runs without knowledge and with the expert statement on the two published '
             'datasets, with 120 characters of context.')
    if len(other):
        terms = sorted({t.split('@')[0] for s in other['recall_terms'] for t in s.split(';')})
        groups = other.groupby(['dataset', 'llm']).size()
        if len(groups) == 1:
            (ds, llm), n = next(iter(groups.items()))
            who = f"The {['', 'one', 'two', 'three', 'four'][n]} {NAME[ds]} run{'s' if n > 1 else ''} of {LLM[llm]}"
        else:
            who = 'The runs ' + '; '.join(run_name(*k) for k in map(tuple, other[['dataset', 'condition', 'llm',
                                                                                   'rep']].values))
        scope += (f" {who} with the misleading statement contained matches of the same kind "
                  f"({', '.join(repr(t) for t in terms)}).")
    text(d, scope)
    for key, hits in sorted(recall.items(), key=lambda kv: keys.index(tuple(kv[0].split('/')))):
        if not hits:
            continue
        ds, cond, llm, rep = key.split('/')
        group_label(d, f'{LABEL[ds]}, {COND[cond]}, {LLM[llm]}, replicate {rep[3:]}')
        for h in hits:
            verbatim(d, f"Turn {h['turn']} ({h['term']}): …{clean_context(h['context'], h['term'])}…")

    heading(d, 'S3.3 Final reports', 3)
    text(d, 'The structured report submitted by the agent with its final model (verbatim).')
    for key in keys:
        res = json.loads((run_dir(*key) / 'results.json').read_text(encoding='utf-8'))
        rep = (res.get('final_model') or {}).get('report') or {}
        group_label(d, f"{LABEL[key[0]]}, {COND[key[1]]}, {LLM[key[2]]}, replicate {key[3][3:]}")
        for k in REPORT_FIELDS:
            if rep.get(k):
                labeled(d, f"{k.replace('_', ' ').capitalize()}. ", rep[k], small=True)


def main():
    """All three sections in one document on the supplement template; the file is saved only when every section was
    built (each section is still attempted, so that all problems are reported at once)."""
    OUT.mkdir(exist_ok=True)
    d = base_document('supplement')
    failed = []
    for build in (s1, s2, s3):
        try:
            build(d)
        except SystemExit as e:
            failed.append(f'Supplementary Material {build.__name__.upper()}: {e}')
    if failed:
        raise SystemExit(f'{FILE} NOT BUILT (the previous file is unchanged):\n' + '\n'.join(failed))
    set_properties(d, f'Supplementary Material: {TITLE}')
    save_docx(d, OUT / FILE)
    print('wrote', OUT / FILE)
    stale = [name for name in OLD_FILES if (OUT / name).exists()]
    if stale:
        print(f'WARNING: {", ".join(stale)} in {OUT} are from the earlier layout with one file per section and are '
              f'not updated; {FILE} replaces them.')


if __name__ == '__main__':
    main()
